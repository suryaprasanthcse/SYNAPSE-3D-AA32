"""Converts the Phase 3 markdown deliverables to styled .docx files.

Usage (from the project root):
    python Phase3_Submission/md_to_docx.py [output_dir]

Writes one .docx per .md next to the sources, and copies them to output_dir when given
(the team uses the Desktop). Handles the markdown subset used by the deliverables:
headings, paragraphs, bullets, tables, block quotes, horizontal rules, **bold** and `code`.
"""

from __future__ import annotations

import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

HERE = Path(__file__).resolve().parent
NAVY = RGBColor(0x14, 0x24, 0x3D)
TEAL = RGBColor(0x0E, 0x6E, 0x7E)
GREY = RGBColor(0x55, 0x5E, 0x6B)
HEADER_FILL = "14243D"
BAND_FILL = "EEF3F7"
TEXT_WIDTH_CM = 16.0


def shade(cell, hex_fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_borders(table, color: str = "8A96A3", size: int = 4) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    table._tbl.tblPr.append(borders)


def add_runs(paragraph, text: str, size: float | None = None, italic: bool = False) -> None:
    """Renders **bold**, `code` and *italic* spans."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|(?<!\*)\*[^*]+\*(?!\*))", text):
        if not part:
            continue
        if part.startswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt((size or 10.5) - 1)
        elif part.startswith("*"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if size and run.font.size is None:
            run.font.size = Pt(size)
        if italic:
            run.italic = True


def configure(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color, before, after in (
        ("Heading 1", 19, NAVY, 0, 8),
        ("Heading 2", 13.5, TEAL, 14, 5),
        ("Heading 3", 11.5, NAVY, 10, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin = section.right_margin = Cm(2.5)
    section.top_margin = section.bottom_margin = Cm(2.2)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def footer(doc: Document, text: str) -> None:
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(f"{text}    ·    Page ")
    run.font.size = Pt(8)
    run.font.color.rgb = GREY
    add_page_number(p)
    for r in p.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = GREY


def horizontal_rule(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("val", "single"), ("sz", "6"), ("space", "2"), ("color", "0E6E7E")):
        bottom.set(qn(f"w:{k}"), v)
    pbdr.append(bottom)
    p._p.get_or_add_pPr().append(pbdr)


def build_table(doc: Document, rows: list[list[str]]) -> None:
    header, body = rows[0], rows[1:]
    cols = len(header)
    width = TEXT_WIDTH_CM / cols
    # A leading narrow column (like "#" or "Slot") reads better tight.
    widths = [width] * cols
    if cols > 2 and len(header[0]) <= 6:
        widths[0] = 2.0
        rest = (TEXT_WIDTH_CM - 2.0) / (cols - 1)
        widths[1:] = [rest] * (cols - 1)

    table = doc.add_table(rows=1, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False

    blank_header = not any(h.strip() for h in header)
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.width = Cm(widths[i])
        if not blank_header:
            shade(cell, HEADER_FILL)
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.bold = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for r, values in enumerate(body):
        cells = table.add_row().cells
        for i, text in enumerate(values[:cols]):
            cell = cells[i]
            cell.width = Cm(widths[i])
            add_runs(cell.paragraphs[0], text, size=9)
            for run in cell.paragraphs[0].runs:
                if run.font.size is None:
                    run.font.size = Pt(9)
            if r % 2 == 1:
                shade(cell, BAND_FILL)

    if blank_header:
        table._tbl.remove(table.rows[0]._tr)
    set_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def convert(md_path: Path) -> Path:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure(doc)
    footer(doc, "LUMINARA · Team Nandha - ML · Funobotz Phase 3")

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if set(stripped) <= {"-", " "} and stripped.startswith("---"):
            horizontal_rule(doc)
            i += 1
            continue

        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            heading = doc.add_heading("", level=min(level, 3))
            add_runs(heading, stripped[level:].strip())
            for run in heading.runs:          # keep the style's colour and weight
                run.font.color.rgb = heading.style.font.color.rgb
                run.bold = True
            i += 1
            continue

        # Table: a header row followed by a separator row of dashes.
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:|-]+\|\s*$", lines[i + 1]):
            rows = [split_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i]))
                i += 1
            build_table(doc, rows)
            continue

        if stripped.startswith(">"):
            p = doc.add_paragraph()
            add_runs(p, stripped.lstrip("> ").strip(), italic=True)
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_before = Pt(6)
            for run in p.runs:
                run.font.color.rgb = TEAL
            pbdr = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            for k, v in (("val", "single"), ("sz", "18"), ("space", "8"), ("color", "0E6E7E")):
                left.set(qn(f"w:{k}"), v)
            pbdr.append(left)
            p._p.get_or_add_pPr().append(pbdr)
            i += 1
            continue

        if re.match(r"^\s*[-*]\s+", line):
            indent = len(line) - len(line.lstrip())
            style = "List Bullet 2" if indent >= 2 else "List Bullet"
            p = doc.add_paragraph(style=style)
            add_runs(p, re.sub(r"^\s*[-*]\s+", "", line))
            p.paragraph_format.space_after = Pt(3)
            i += 1
            continue

        if re.match(r"^\s*\d+\.\s+", line):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\s*\d+\.\s+", "", line))
            p.paragraph_format.space_after = Pt(3)
            i += 1
            continue

        # Plain paragraph: join wrapped lines until a blank or a new block starts.
        buffer = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if not nxt or nxt.startswith(("#", "|", ">", "-", "*")) or re.match(r"^\d+\.\s", nxt):
                break
            buffer.append(nxt)
            i += 1
        add_runs(doc.add_paragraph(), " ".join(buffer))

    out = md_path.with_suffix(".docx")
    doc.core_properties.title = md_path.stem.replace("_", " ")
    doc.core_properties.author = "Team Nandha - ML"
    doc.save(out)
    return out


def main() -> None:
    destination = Path(sys.argv[1]).expanduser() if len(sys.argv) > 1 else None
    if destination:
        destination.mkdir(parents=True, exist_ok=True)

    for md in sorted(HERE.glob("*.md")):
        docx = convert(md)
        note = ""
        if destination:
            shutil.copy2(docx, destination / docx.name)
            note = f" -> {destination / docx.name}"
        print(f"{md.name} -> {docx.name}{note}")


if __name__ == "__main__":
    main()
